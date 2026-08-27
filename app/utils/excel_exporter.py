import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from app.models import LPLPO

def export_lplpo_to_excel(lplpo_id):
    """
    Mengekspor data LPLPO ke file Excel (.xlsx) dengan formatting rapi sesuai standar Kemenkes/Dinkes.
    """
    lplpo = LPLPO.query.get(lplpo_id)
    if not lplpo:
        raise ValueError(f"LPLPO dengan ID {lplpo_id} tidak ditemukan.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"LPLPO {lplpo.nama_bulan} {lplpo.tahun}"

    # Custom styling
    font_title = Font(name='Arial', size=14, bold=True)
    font_subtitle = Font(name='Arial', size=11, bold=True)
    font_header = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    font_data = Font(name='Arial', size=10)
    font_total = Font(name='Arial', size=10, bold=True)

    fill_header = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')  # Deep Blue
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Header Judul
    ws.merge_cells('A1:L1')
    ws['A1'] = "LAPORAN PEMAKAIAN DAN LEMBAR PERMINTAAN OBAT (LPLPO)"
    ws['A1'].font = font_title
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A2:L2')
    ws['A2'] = f"PUSKESMAS: {lplpo.nama_puskesmas.upper()} | PERIODE: {lplpo.nama_bulan.upper()} {lplpo.tahun}"
    ws['A2'].font = font_subtitle
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    # Header Tabel
    headers = [
        "No", "Kode Obat", "Nama Obat", "Satuan",
        "Stok Awal", "Penerimaan", "Persediaan", "Pemakaian",
        "Sisa Stok", "Stok Optimum", "Permintaan", "Pemberian"
    ]

    header_row = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[header_row].height = 28

    # Data Rows
    start_row = 5
    items = sorted(lplpo.items, key=lambda x: x.obat.nama_obat)

    total_stok_awal = 0
    total_penerimaan = 0
    total_persediaan = 0
    total_pemakaian = 0
    total_sisa_stok = 0

    for idx, item in enumerate(items, 1):
        current_row = start_row + idx - 1
        
        row_data = [
            idx,
            item.obat.kode_obat,
            item.obat.nama_obat,
            item.obat.satuan,
            item.stok_awal,
            item.penerimaan,
            item.persediaan,
            item.pemakaian,
            item.sisa_stok,
            item.stok_optimum,
            item.permintaan,
            item.pemberian
        ]

        total_stok_awal += item.stok_awal
        total_penerimaan += item.penerimaan
        total_persediaan += item.persediaan
        total_pemakaian += item.pemakaian
        total_sisa_stok += item.sisa_stok

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.border = thin_border

            if col_idx in [1, 2, 4]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'

            if current_row % 2 == 0:
                cell.fill = fill_zebra

        ws.row_dimensions[current_row].height = 20

    # Total Row
    total_row = start_row + len(items)
    ws.cell(row=total_row, column=1, value="TOTAL")
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.cell(row=total_row, column=1).font = font_total
    ws.cell(row=total_row, column=1).alignment = Alignment(horizontal='center', vertical='center')

    totals = {
        5: total_stok_awal,
        6: total_penerimaan,
        7: total_persediaan,
        8: total_pemakaian,
        9: total_sisa_stok
    }

    for col_idx in range(1, 13):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.font = font_total
        cell.border = thin_border
        if col_idx in totals:
            cell.value = totals[col_idx]
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.number_format = '#,##0'

    ws.row_dimensions[total_row].height = 24

    # Adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.column_dimensions['C'].width = 32  # Nama Obat wider

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
